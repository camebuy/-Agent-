"""Create and revise local deliverables without changing the user's source files."""
from __future__ import annotations

import json
import os
import re
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from src.core.system.path_manager import PathManager
from src.tools.database_tool import get_file_path


_TEXT_SUFFIXES = {".txt", ".md", ".docx", ".pdf"}
_ALL_SUFFIXES = _TEXT_SUFFIXES | {".xlsx"}
_MAX_SOURCE_BYTES = 20 * 1024 * 1024
_MAX_CONTENT_CHARS = 100_000
_MAX_ROWS = 5_000


def _name(value: str, suffixes: set[str]) -> str:
    name = str(value or "").strip()
    if (not name or len(name) > 120 or name in {".", ".."}
            or Path(name).name != name or "\\" in name
            or re.search(r'[<>:"/|?*\x00-\x1f]', name)
            or name.rstrip(" .") != name
            or name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}):
        raise ValueError("只接受普通文件名，不接受路径或 Windows 保留名称。")
    if Path(name).suffix.lower() not in suffixes:
        raise ValueError(f"文件后缀必须是：{', '.join(sorted(suffixes))}。")
    return name


def _output(name: str, suffixes: set[str] = _ALL_SUFFIXES) -> Path:
    return PathManager.DOCUMENT_OUTPUT_DIR / _name(name, suffixes)


@contextmanager
def _new_output(path: Path) -> Iterator[object]:
    """Exclusive creation: a generated version never overwrites an earlier one."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as stream:
        try:
            yield stream
        except BaseException:
            stream.close()
            path.unlink(missing_ok=True)
            raise


def _source(source_name: str, source_file_id: str, suffixes: set[str]) -> Path:
    if bool(source_name) == bool(source_file_id):
        raise ValueError("source_name 和 source_file_id 必须且只能提供一个。")
    if source_file_id:
        from src.core.system.runtime_settings import load_runtime_settings
        if not load_runtime_settings()["tool_permissions"].get("user_files", True):
            raise PermissionError("读取用户原件需要启用‘读取和管理我的文件’权限。")
        path = get_file_path(source_file_id)
    else:
        path = _output(source_name, suffixes)
    if path.suffix.lower() not in suffixes or not path.is_file() or path.is_symlink():
        raise ValueError("来源文件不存在或格式不支持。")
    if path.stat().st_size > _MAX_SOURCE_BYTES:
        raise ValueError("来源文件超过 20 MiB，停止处理。")
    return path


def _result(path: Path, action: str) -> str:
    return json.dumps({"status": action, "name": path.name, "path": str(path),
                       "size_bytes": path.stat().st_size}, ensure_ascii=False)


def document_create_text(name: str, content: str) -> str:
    """Create a new TXT, Markdown, DOCX or PDF file from supplied text."""
    path = _output(name, _TEXT_SUFFIXES)
    content = str(content or "")
    if not content.strip() or len(content) > _MAX_CONTENT_CHARS:
        raise ValueError("正文须为 1 至 100000 个字符。")
    suffix = path.suffix.lower()
    with _new_output(path) as stream:
        if suffix in {".txt", ".md"}:
            stream.write(content.encode("utf-8"))
        elif suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("缺少 python-docx，请在当前 Python 环境安装 requirements.txt。") from exc
            doc = Document()
            for line in content.splitlines():
                doc.add_paragraph(line)
            doc.save(stream)
        else:
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.lib.styles import ParagraphStyle
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.cidfonts import UnicodeCIDFont
                from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
                from xml.sax.saxutils import escape
            except ImportError as exc:
                try:
                    import fitz
                except ImportError:
                    raise RuntimeError("生成 PDF 需要 reportlab 或 PyMuPDF；请安装 requirements.txt。") from exc
                font_file = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "simhei.ttf"
                if not font_file.is_file():
                    raise RuntimeError("缺少可用于中文 PDF 的 SimHei 字体，请安装 reportlab。") from exc
                font = fitz.Font(fontfile=str(font_file))
                pdf = fitz.open()
                page = pdf.new_page(width=595, height=842)
                page.insert_font(fontname="AACWorkHei", fontfile=str(font_file))
                y = 60
                for line in content.splitlines():
                    segments = []
                    segment = ""
                    for char in line:
                        proposed = segment + char
                        if segment and font.text_length(proposed, fontsize=11) > 499:
                            segments.append(segment)
                            segment = char
                        else:
                            segment = proposed
                    segments.append(segment)
                    for part in segments:
                        if y > 794:
                            page = pdf.new_page(width=595, height=842)
                            page.insert_font(fontname="AACWorkHei", fontfile=str(font_file))
                            y = 60
                        if part:
                            page.insert_text((48, y), part, fontsize=11, fontname="AACWorkHei")
                        y += 17
                stream.write(pdf.tobytes())
                pdf.close()
            else:
                pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
                style = ParagraphStyle("AACWorkBody", fontName="STSong-Light", fontSize=11, leading=17)
                story = [Paragraph(escape(line) or "&#160;", style) if line else Spacer(1, 17)
                         for line in content.splitlines()]
                SimpleDocTemplate(stream, pagesize=A4, leftMargin=48, rightMargin=48,
                                  topMargin=48, bottomMargin=48).build(story)
    return _result(path, "CREATED")


def document_create_spreadsheet(name: str, headers: list[str], rows: list[list[str]]) -> str:
    """Create a simple XLSX table; cell values are text, not executable formulas."""
    path = _output(name, {".xlsx"})
    headers = [str(value) for value in (headers or [])]
    rows = rows or []
    if not headers or len(headers) > 50 or len(rows) > _MAX_ROWS:
        raise ValueError("表格需要 1–50 列且最多 5000 行。")
    if any(len(value) > 500 for value in headers):
        raise ValueError("表头文字过长。")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请在当前 Python 环境安装 requirements.txt。") from exc
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    for index, value in enumerate(headers, 1):
        cell = sheet.cell(1, index, value)
        cell.data_type = "s"
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, 2):
        if not isinstance(row, list) or len(row) > len(headers):
            raise ValueError(f"第 {row_index} 行的列数超过表头。")
        for column, value in enumerate(row, 1):
            text = str(value if value is not None else "")
            if len(text) > 32_000:
                raise ValueError("单元格文字过长。")
            cell = sheet.cell(row_index, column, text)
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    with _new_output(path) as stream:
        book.save(stream)
    return _result(path, "CREATED")


def _docx_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def document_revise_text(output_name: str, find_text: str, replacement: str,
                         source_name: str = "", source_file_id: str = "") -> str:
    """Replace exact text in TXT/MD/DOCX and save a separate version."""
    source = _source(source_name, source_file_id, {".txt", ".md", ".docx"})
    output = _output(output_name, {source.suffix.lower()})
    needle = str(find_text or "")
    replacement = str(replacement or "")
    if not needle or len(needle) > 10_000 or len(replacement) > 20_000:
        raise ValueError("需要明确的查找文字，替换文字不能超过 20000 字符。")
    if source.resolve() == output.resolve():
        raise ValueError("修订版必须使用新的文件名。")
    if source.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，请安装 requirements.txt。") from exc
        doc = Document(source)
        count = 0
        for paragraph in _docx_paragraphs(doc):
            visible_count = paragraph.text.count(needle)
            run_count = sum(run.text.count(needle) for run in paragraph.runs)
            if visible_count != run_count:
                raise ValueError("待替换文字跨越 Word 格式片段；已停止，避免破坏排版。")
            for run in paragraph.runs:
                count += run.text.count(needle)
                run.text = run.text.replace(needle, replacement)
        if not count:
            raise ValueError("未找到完全匹配的文字，没有生成修订版。")
        with _new_output(output) as stream:
            doc.save(stream)
    else:
        original = source.read_text(encoding="utf-8")
        count = original.count(needle)
        if not count:
            raise ValueError("未找到完全匹配的文字，没有生成修订版。")
        revised = original.replace(needle, replacement)
        if len(revised) > _MAX_CONTENT_CHARS:
            raise ValueError("修订后超过 100000 字符。")
        with _new_output(output) as stream:
            stream.write(revised.encode("utf-8"))
    result = json.loads(_result(output, "REVISED"))
    result["replacements"] = count
    return json.dumps(result, ensure_ascii=False)


def document_append_text(output_name: str, text: str,
                         source_name: str = "", source_file_id: str = "") -> str:
    """Append text to TXT/MD/DOCX, preserving the original as a separate version."""
    source = _source(source_name, source_file_id, {".txt", ".md", ".docx"})
    output = _output(output_name, {source.suffix.lower()})
    text = str(text or "")
    if not text.strip() or len(text) > _MAX_CONTENT_CHARS:
        raise ValueError("追加文字须为 1–100000 字符。")
    if source.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，请安装 requirements.txt。") from exc
        doc = Document(source)
        for line in text.splitlines():
            doc.add_paragraph(line)
        with _new_output(output) as stream:
            doc.save(stream)
    else:
        old = source.read_text(encoding="utf-8")
        if len(old) + len(text) > _MAX_CONTENT_CHARS:
            raise ValueError("追加后文件超过 100000 字符。")
        separator = "" if not old or old.endswith("\n") else "\n"
        with _new_output(output) as stream:
            stream.write((old + separator + text).encode("utf-8"))
    return _result(output, "REVISED")


def document_append_spreadsheet_rows(output_name: str, rows: list[list[str]],
                                     source_name: str = "", source_file_id: str = "") -> str:
    """Append rows to a selected workbook and save a separate XLSX version."""
    source = _source(source_name, source_file_id, {".xlsx"})
    output = _output(output_name, {".xlsx"})
    rows = rows or []
    if not rows or len(rows) > _MAX_ROWS:
        raise ValueError("需要 1–5000 行新数据。")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请安装 requirements.txt。") from exc
    book = load_workbook(source, read_only=False, keep_links=False)
    sheet = book.active
    columns = sheet.max_column
    if columns < 1 or columns > 50:
        raise ValueError("来源表格的列数不在支持范围内。")
    for row in rows:
        if not isinstance(row, list) or len(row) > columns:
            raise ValueError("新数据列数超过来源表格。")
        cells = [str(value if value is not None else "") for value in row]
        if any(len(value) > 32_000 for value in cells):
            raise ValueError("单元格文字过长。")
        sheet.append(cells)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    with _new_output(output) as stream:
        book.save(stream)
    book.close()
    result = json.loads(_result(output, "REVISED"))
    result["appended_rows"] = len(rows)
    return json.dumps(result, ensure_ascii=False)


def document_update_spreadsheet_cells(output_name: str, updates: list[dict[str, str]],
                                      source_name: str = "", source_file_id: str = "",
                                      sheet_name: str = "") -> str:
    """Change exact XLSX cells only when their observed old values still match."""
    source = _source(source_name, source_file_id, {".xlsx"})
    output = _output(output_name, {".xlsx"})
    if not isinstance(updates, list) or not updates or len(updates) > 100:
        raise ValueError("需要 1–100 个单元格修改。")
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请安装 requirements.txt。") from exc
    book = load_workbook(source, read_only=False, keep_links=False)
    try:
        if sheet_name and sheet_name not in book.sheetnames:
            raise ValueError("找不到指定工作表。")
        sheet = book[sheet_name] if sheet_name else book.active
        checked = []
        seen = set()
        for item in updates:
            if not isinstance(item, dict) or not {"cell", "expected", "value"} <= item.keys():
                raise ValueError("每项修改都需提供 cell、expected 和 value。")
            address = str(item["cell"] or "").strip().upper()
            try:
                letters, row = coordinate_from_string(address)
                column = column_index_from_string(letters)
            except ValueError as exc:
                raise ValueError(f"无效单元格地址：{address}") from exc
            if not (1 <= column <= 50 and 1 <= row <= 5000) or address in seen:
                raise ValueError("单元格超出 A1:AX5000 或被重复修改。")
            seen.add(address)
            current = sheet[address].value
            if str(current if current is not None else "") != str(item["expected"]):
                raise ValueError(f"{address} 的现有内容与 expected 不同，已停止修改。")
            value = str(item["value"] if item["value"] is not None else "")
            if len(value) > 32_000:
                raise ValueError("单元格文字过长。")
            checked.append((address, value))
        for address, value in checked:
            sheet[address] = value
            sheet[address].data_type = "s"
        with _new_output(output) as stream:
            book.save(stream)
    finally:
        book.close()
    result = json.loads(_result(output, "REVISED"))
    result["updated_cells"] = len(checked)
    return json.dumps(result, ensure_ascii=False)


def document_list_outputs() -> str:
    """List generated deliverables; source user files remain in their own catalog."""
    root = PathManager.DOCUMENT_OUTPUT_DIR
    if not root.is_dir():
        return json.dumps({"files": []}, ensure_ascii=False)
    files = sorted((p for p in root.iterdir() if p.is_file() and not p.is_symlink()
                    and p.suffix.lower() in _ALL_SUFFIXES), key=lambda p: p.stat().st_mtime, reverse=True)
    return json.dumps({"files": [{"name": p.name, "path": str(p), "size_bytes": p.stat().st_size}
                                  for p in files[:100]]}, ensure_ascii=False)


def document_read_output(name: str, max_chars: int = 20_000) -> str:
    """Read a generated file so the Agent can inspect the delivered content."""
    path = _source(name, "", _ALL_SUFFIXES)
    limit = max(1, min(int(max_chars), 20_000))
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        value = path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        from docx import Document
        value = "\n".join(p.text for p in _docx_paragraphs(Document(path)))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            value = "\n".join("\t".join(str(cell if cell is not None else "") for cell in row)
                              for row in book.active.iter_rows(max_row=200, max_col=50, values_only=True))
        finally:
            book.close()
    else:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("缺少 pypdf，请安装 requirements.txt。") from exc
        value = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages[:30])
    return json.dumps({"name": path.name, "text": value[:limit], "truncated": len(value) > limit},
                      ensure_ascii=False)
